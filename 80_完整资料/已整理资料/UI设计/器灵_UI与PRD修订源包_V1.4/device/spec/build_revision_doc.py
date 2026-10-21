from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
R=Path(__file__).parent
p=R/'器灵_非触摸显示屏_补充修订.md'
d=Document();s=d.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=s.bottom_margin=Cm(1.7);s.left_margin=s.right_margin=Cm(1.9)
for name in ['Normal','Title','Heading 1','Heading 2','List Bullet']:
 st=d.styles[name];st.font.name='Noto Sans CJK SC';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
 st.font.size=Pt({'Normal':10,'Title':22,'Heading 1':15,'Heading 2':11,'List Bullet':10}[name]);st.font.color.rgb=RGBColor.from_string('000000')
 st.paragraph_format.line_spacing=1.08;st.paragraph_format.space_after=Pt(4)
 if name.startswith('Heading'):st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(9)
for line in p.read_text().splitlines():
 if not line:continue
 if line.startswith('# '):q=d.add_paragraph(line[2:],'Title')
 elif line.startswith('## '):q=d.add_paragraph(line[3:],'Heading 1')
 elif line.startswith('### '):q=d.add_paragraph(line[4:],'Heading 2')
 elif line.startswith('- '):q=d.add_paragraph(line[2:],'List Bullet')
 else:q=d.add_paragraph(line)
 q.paragraph_format.widow_control=True
f=s.footer.paragraphs[0];f.alignment=2
r=f.add_run('器灵 · 非触摸显示屏修订 R2 / V1.4  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
for node in list(d.styles.element.iter(qn('w:pBdr'))) + list(d.element.iter(qn('w:pBdr'))):
 node.getparent().remove(node)
d.core_properties.title='器灵 · 非触摸显示屏补充修订';d.core_properties.author='dot'
d.save(R/'器灵_非触摸显示屏_补充修订.docx')
